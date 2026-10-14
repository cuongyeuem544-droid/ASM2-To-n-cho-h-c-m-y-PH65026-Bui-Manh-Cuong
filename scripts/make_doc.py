import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Thiết lập lề 2 cm (khoảng 0.8 inch)
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Định dạng chung
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(10.5)
font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

def add_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

def add_single_line_item(num, bold_prefix, text_content):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25)
    
    # Bullet number
    run_num = p.add_run(f"{num}. ")
    run_num.font.name = 'Arial'
    run_num.font.size = Pt(10.5)
    run_num.font.bold = True
    run_num.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    # Bold prefix
    run_prefix = p.add_run(bold_prefix)
    run_prefix.font.name = 'Arial'
    run_prefix.font.size = Pt(10.5)
    run_prefix.font.bold = True
    run_prefix.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)
    
    # Content
    run_text = p.add_run(f": {text_content}")
    run_text.font.name = 'Arial'
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

add_title("H4CARE / PHARMATRUST - TỔNG HỢP THAY ĐỔI TRONG 2 NGÀY")
add_subtitle("Khoảng thời gian: 11/10/2026 - 13/10/2026 | Tác giả: NGUYENDUCANHPH65820 | Hệ thống: H4CARE / PharmaTrust")

add_heading("I. NHỮNG NỘI DUNG ĐÃ SỬA (MỖI MỤC ĐÚNG 1 DÒNG DUY NHẤT)")

sua_items = [
    ("Sửa lỗi liên kết tồn kho thuốc", "Kết nối thuốc đã duyệt nhập khẩu với số lô và kho vật lý, sửa dứt điểm lỗi HTTP 500 khi mở danh sách lô."),
    ("Sửa thẩm định mở bán sản phẩm", "Đánh giá chuẩn điều kiện thương mại theo từng SKU và đồng bộ hóa thời gian GMT+7 Việt Nam cho toàn hệ thống."),
    ("Tối ưu tốc độ cuộn Storefront", "Tắt hiệu ứng cuộn quán tính nhân tạo của Lenis để khôi phục độ mượt và phản hồi cuộn tự nhiên của trình duyệt."),
    ("Sửa lỗi tràn giao diện bảng Admin", "Tích hợp thanh phân trang PaginationBar đồng bộ trên toàn bộ 17 trang quản trị để chống vỡ khung dọc."),
    ("Sửa lỗi kết nối máy chủ và lỗi 405", "Bổ sung route alias @router.patch('/orders/{order_id}/status') song song với /fulfillment khi bàn giao đơn kho."),
    ("Sửa lỗi 500 khi xem Phiếu Soạn (Pick Slip)", "Chuẩn hóa tên trường model ORM OrderItem (product_id, price, subtotal), ProductSku (uom) và thuật toán FEFO."),
    ("Sửa lỗi lịch sử đơn hàng khách bị rỗng", "Bỏ cơ chế so sánh theo số điện thoại, sửa lỗi nuốt exception mạng thành mảng rỗng, chuyển sang xác thực Bearer token."),
    ("Sửa checkout nuốt token sai", "Thay thế get_optional_current_user bằng strict_optional_current_user để từ chối token hỏng/hết hạn bằng HTTP 401."),
    ("Sửa lỗ hổng xung đột Idempotency Key", "Bắt buộc kiểm tra quyền sở hữu order.user_id == req_user_id khi gửi lại đơn, trả mã HTTP 409 nếu khác tài khoản."),
    ("Sửa URL máy chủ gắn cứng trên Storefront", "Thay thế toàn bộ địa chỉ http://localhost:8000 bằng biến môi trường process.env.NEXT_PUBLIC_API_URL."),
    ("Sửa logic phân quyền chéo kho", "Bắt buộc lọc warehouse_id và chặn triệt để việc Quản lý kho xem hoặc can thiệp dữ liệu của chi nhánh kho khác."),
]

for idx, (prefix, text) in enumerate(sua_items, 1):
    add_single_line_item(idx, prefix, text)

add_heading("II. NHỮNG NỘI DUNG ĐÃ BỔ SUNG (MỖI MỤC ĐÚNG 1 DÒNG DUY NHẤT)")

bo_sung_items = [
    ("Bổ sung modal chi tiết tồn kho theo từng kho", "Thêm cửa sổ tra cứu trực quan số lượng tồn thực tế, số lượng giữ chỗ và tồn khả dụng tại trang Kho hàng."),
    ("Bổ sung hệ thống duyệt phân loại thuốc", "Thêm trang danh mục và hàng đợi duyệt phân loại tự động (CategoriesPage, ClassificationQueuePage)."),
    ("Bổ sung cơ chế phân quyền Đa kho (RBAC)", "Tạo migration 0005 thêm các vai trò WAREHOUSE_MANAGER, WAREHOUSE_STAFF và liên kết UserWarehouseMembership."),
    ("Bổ sung Cổng thông tin Quản lý kho", "Xây dựng menu, layout riêng và các trang: Tổng quan, Tồn kho, Lô thuốc, Phiếu nhập, Điều chuyển, Kiểm kê, Nhân viên, Nhật ký."),
    ("Bổ sung Cổng vận hành Nhân viên kho", "Xây dựng giao diện công việc hôm nay, danh sách soạn hàng/đóng gói, nhận hàng và lịch sử tác vụ."),
    ("Bổ sung Module quản lý kho chi nhánh riêng", "Thêm trang Sản phẩm trong kho (WarehouseProductsPage), Cảnh báo tồn/hạn (WarehouseAlertsPage), Nhập kho file (WarehouseImportsPage)."),
    ("Bổ sung trang Quản lý tài khoản bảo mật", "Admin quản trị toàn hệ thống, Quản lý kho chỉ quản lý nhân viên kho mình, mật khẩu tạm chỉ hiện 1 lần duy nhất (UsersPage)."),
    ("Bổ sung cơ chế bắt buộc đổi mật khẩu", "Thêm trường must_change_password, modal ForceChangePasswordModal chặn thao tác và hủy phiên cũ (token_version)."),
    ("Bổ sung API lịch sử đơn hàng theo tài khoản", "Tạo endpoint GET /api/v1/store/orders/me xác thực qua JWT token, eager-load OrderItem, chống IDOR 100%."),
    ("Bổ sung luồng tra cứu đơn cho khách vãng lai", "Tạo endpoint POST /api/v1/store/orders/guest-lookup yêu cầu khớp cả mã đơn và SĐT, tự động che bảo mật số điện thoại."),
    ("Bổ sung 4 trạng thái giao diện lịch sử đơn hàng", "Thiết kế đủ 4 trạng thái loading (đang tải), error (báo lỗi kèm nút Thử lại), success-data (có đơn), và success-empty (chưa có đơn)."),
    ("Bổ sung công cụ di chuyển đơn hàng cũ", "Viết script link_legacy_orders.py với cờ an toàn --dry-run và --commit để ghép 19 đơn hàng cũ vào tài khoản người dùng tương ứng."),
    ("Bổ sung bộ kiểm thử tự động toàn diện", "Xây dựng 4 bộ test suite mới với hàng chục test cases bao phủ RBAC, quản lý kho chi nhánh, bảo mật tài khoản và lịch sử đơn hàng."),
]

for idx, (prefix, text) in enumerate(bo_sung_items, 1):
    add_single_line_item(idx, prefix, text)

# Lưu tại cả 2 thư mục để người dùng mở từ đâu cũng thấy
path1 = r"C:\Users\admin\Desktop\TOAN HOC CHO MAY\Bao_cao_tong_hop_thay_doi_2_ngay_qua_H4CARE.docx"
path2 = r"C:\Users\admin\Desktop\TOAN HOC CHO MAY\du-an-website-ban-thuoc-\Bao_cao_tong_hop_thay_doi_2_ngay_qua_H4CARE.docx"

doc.save(path1)
doc.save(path2)
print("SUCCESS: Đã tạo file thành công!")
print("Path 1:", path1)
print("Path 2:", path2)
